"""Customer import uses explicit mappings, encrypted staging and tenant-serialized inserts."""
import csv
from datetime import date, datetime
from io import BytesIO, StringIO
from pathlib import Path
import unicodedata
from zipfile import BadZipFile, ZipFile
from xml.etree.ElementTree import ParseError
from defusedxml.common import DefusedXmlException
from lxml.etree import XMLSyntaxError
from openpyxl.utils.exceptions import InvalidFileException

from django.core.exceptions import ValidationError
from django.db import transaction
from django.utils import timezone

from communications.phone import whatsapp_number
from communications.platform_events import after_commit_notice, suppress_platform_events
from core.audit import append_audit
from core.crypto import decrypt_json
from scheduling.customer_identity import contact_values
from scheduling.models import Customer, CustomerImportBatch
from tenants.models import Tenant

MAX_BYTES = 5 * 1024 * 1024
MAX_ROWS = 5000
MAX_COLUMNS = 80
FIELDS = {"name": "Nome completo", "phone": "Telefone / WhatsApp com DDD",
          "email": "E-mail (opcional)", "birth_date": "Nascimento (opcional)",
          "preferences": "Preferências / observações (opcional)"}


def text(value):
    if value is None:
        return ""
    if isinstance(value, (datetime, date)):
        return value.date().isoformat() if isinstance(value, datetime) else value.isoformat()
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def read_source(content, filename):
    if not content or len(content) > MAX_BYTES:
        raise ValidationError("Envie um CSV ou XLSX de até 5 MB.")
    suffix = Path(filename).suffix.lower()
    try:
        if suffix == ".csv":
            decoded = None
            for encoding in ("utf-8-sig", "cp1252"):
                try:
                    decoded = content.decode(encoding)
                    break
                except UnicodeDecodeError:
                    continue
            if decoded is None or "\x00" in decoded:
                raise ValueError
            try:
                dialect = csv.Sniffer().sniff(decoded[:8192], delimiters=";,\t|")
                source = csv.reader(StringIO(decoded), dialect)
            except csv.Error:
                source = csv.reader(StringIO(decoded), delimiter=";")
            rows = []
            for row in source:
                if len(rows) >= MAX_ROWS + 1:
                    raise ValidationError("Divida a base em arquivos de até 5.000 clientes.")
                if any(text(cell) for cell in row):
                    rows.append(row)
        elif suffix == ".xlsx":
            with ZipFile(BytesIO(content)) as archive:
                members = archive.infolist()
                if len(members) > 2000 or sum(item.file_size for item in members) > 32 * 1024 * 1024:
                    raise ValidationError("A planilha descompactada excede o limite de segurança.")
            from openpyxl import load_workbook
            workbook = load_workbook(BytesIO(content), read_only=True, data_only=False, keep_links=False)
            try:
                sheet = workbook.worksheets[0]
                rows = []
                for scanned, row in enumerate(sheet.iter_rows(), 1):
                    if scanned > MAX_ROWS + 1:
                        raise ValidationError("Remova linhas vazias excedentes ou divida a planilha em até 5.000 clientes.")
                    if len(rows) >= MAX_ROWS + 1:
                        raise ValidationError("Divida a base em arquivos de até 5.000 clientes.")
                    if len(row) > MAX_COLUMNS:
                        raise ValidationError("Use no máximo 80 colunas.")
                    if any(cell.data_type == "f" for cell in row):
                        raise ValidationError("Converta fórmulas em valores antes de importar.")
                    values = [text(cell.value) for cell in row]
                    if any(values):
                        rows.append(values)
            finally:
                workbook.close()
        else:
            raise ValidationError("Use CSV ou XLSX. Exporte backups proprietários pelo sistema de origem.")
    except ValidationError:
        raise
    except (ValueError, OSError, BadZipFile, csv.Error, KeyError, IndexError,
            ParseError, DefusedXmlException, XMLSyntaxError, InvalidFileException):
        raise ValidationError("Não foi possível ler o arquivo. Confira o formato e o cabeçalho.") from None
    if sum(len(text(cell)) for row in rows for cell in row) > MAX_BYTES:
        raise ValidationError("O conteúdo textual da planilha excede 5 MB. Remova colunas desnecessárias ou divida a base.")
    if len(rows) < 2:
        raise ValidationError("O arquivo deve ter cabeçalho e pelo menos um cliente.")
    headers = [text(cell) for cell in rows[0]]
    if not headers or len(headers) > MAX_COLUMNS:
        raise ValidationError("Use de 1 a 80 colunas.")
    cleaned = []
    for row in rows[1:]:
        if len(row) > len(headers) or any(len(text(cell)) > 1000 for cell in row):
            raise ValidationError("Há linhas com mais colunas que o cabeçalho ou campos com mais de 1.000 caracteres.")
        cleaned.append([text(cell) for cell in row] + [""] * (len(headers) - len(row)))
    return headers, cleaned


def suggested_mapping(headers):
    aliases = {"name": {"nome", "nome completo", "cliente", "nome cliente", "name", "nome do cliente"},
               "phone": {"telefone", "celular", "whatsapp", "fone", "phone", "telefone celular"},
               "email": {"email", "e mail"}, "birth_date": {"nascimento", "data nascimento", "data de nascimento"},
               "preferences": {"observacoes", "preferencias", "anotacoes"}}
    normalized = [" ".join("".join(c for c in unicodedata.normalize("NFD", header.lower())
                                 if not unicodedata.combining(c)).replace("-", " ").replace("_", " ").split()) for header in headers]
    return {field: str(next(index for index, header in enumerate(normalized) if header in values))
            for field, values in aliases.items() if any(header in values for header in normalized)}


def preview_rows(tenant, headers, rows, mapping):
    phones, emails = {}, {}
    def remember(phone, email, identity):
        if phone:
            phones.setdefault(phone, set()).add(identity)
        if email:
            emails.setdefault(email.lower(), set()).add(identity)
    for existing in Customer.objects.filter(tenant=tenant).values("pk", "phone", "email"):
        remember(whatsapp_number(existing["phone"]), existing["email"].strip(), existing["pk"])
    output = []
    for number, row in enumerate(rows, 2):
        result = {"row": number, "state": "error", "reason": "", "data": {}}
        def cell(field):
            selected = mapping.get(field, "")
            if selected == "":
                return ""
            index = int(selected)
            if not 0 <= index < len(headers):
                raise ValidationError("Mapeamento de colunas inválido.")
            return row[index]
        try:
            name, phone, email = contact_values(cell("name"), cell("phone"), cell("email"))
            if len(email) > 254:
                raise ValidationError("O e-mail deve ter até 254 caracteres.")
            birth = cell("birth_date")
            parsed_birth = None
            if birth:
                for format_ in ("%Y-%m-%d", "%d/%m/%Y", "%d-%m-%Y"):
                    try:
                        parsed_birth = datetime.strptime(birth, format_).date()
                        break
                    except ValueError:
                        continue
                if parsed_birth is None or parsed_birth > timezone.localdate():
                    raise ValidationError("Data de nascimento inválida.")
            identities = phones.get(phone, set()) | (emails.get(email, set()) if email else set())
            if len(identities) > 1:
                raise ValidationError("Telefone e e-mail vinculados a cadastros diferentes. Revise antes de importar.")
            result["data"] = {"name": name, "phone": phone, "email": email,
                              "birth_date": parsed_birth.isoformat() if parsed_birth else "",
                              "preferences": cell("preferences")}
            if identities:
                result.update(state="skip", reason="Já existente na empresa ou repetido no arquivo; cadastro preservado.")
            else:
                result.update(state="create", reason="Novo cliente.")
                remember(phone, email, -number)
        except (ValidationError, ValueError, TypeError) as exc:
            result["reason"] = " ".join(exc.messages) if isinstance(exc, ValidationError) else "Linha ou mapeamento inválido."
        output.append(result)
    return output


def counts(rows):
    return {state: sum(row["state"] == state for row in rows) for state in ("create", "skip", "error")}


@transaction.atomic
def apply_batch(batch_id, actor, allow_errors=False, expected_mapping=None):
    batch = CustomerImportBatch.objects.select_for_update().select_related("tenant").get(pk=batch_id)
    if not actor.is_superuser or not actor.is_active or getattr(actor, "deleted_at", None):
        raise ValidationError("Somente o Master ativo pode importar bases.")
    if batch.state == "applied":
        return batch
    if batch.expires_at <= timezone.now():
        raise ValidationError("A prévia expirou. Envie o arquivo novamente.")
    if expected_mapping is None or batch.mapping != expected_mapping:
        raise ValidationError("O mapeamento mudou. Gere uma nova prévia antes de confirmar.")
    tenant = Tenant.objects.select_for_update().get(pk=batch.tenant_id, deleted_at__isnull=True, archived_at__isnull=True)
    payload = decrypt_json(batch.payload_encrypted)
    rows = preview_rows(tenant, payload["headers"], payload["rows"], batch.mapping)
    summary = counts(rows)
    if summary["error"] and not allow_errors:
        raise ValidationError("Há linhas inválidas. Corrija o arquivo ou autorize explicitamente importar apenas as válidas.")
    objects = []
    for row in rows:
        if row["state"] == "create":
            data = dict(row["data"])
            data["birth_date"] = date.fromisoformat(data["birth_date"]) if data["birth_date"] else None
            objects.append(Customer(tenant=tenant, consent_marketing=False, **data))
    with suppress_platform_events():
        Customer.objects.bulk_create(objects, batch_size=500)
        append_audit(tenant=tenant, user=actor, action="CUSTOMER_BASE_IMPORTED",
                     entity_type="CustomerImportBatch", entity_id=batch.pk, after=summary)
    batch.state = "applied"
    batch.applied_at = timezone.now()
    batch.payload_encrypted = ""
    batch.report = {"counts": summary, "rows": [
        {"row": row["row"], "state": row["state"], "reason": row["reason"]} for row in rows]}
    batch.save(update_fields=["state", "applied_at", "payload_encrypted", "report", "updated_at"])
    after_commit_notice(f"customer-import:{batch.pk}", "Base de clientes importada",
                        f"/master/empresas/{tenant.pk}/ficha/")
    return batch
